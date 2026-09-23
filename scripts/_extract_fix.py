# -*- coding: utf-8 -*-
"""提取脚本 v2：修复标题行在表头前的 bug + 品牌去重"""
import openpyxl, os, re, json

P = r"C:\Users\shenghua\Desktop\新建文件夹"
FILES = ['学习文具.xlsx','体育用品.xlsx','少儿穿戴.xlsx','读物潮玩.xlsx','少儿食品.xlsx',
         '家庭日用.xlsx','米面油品.xlsx','家用电器.xlsx','少儿护理.xlsx']


def is_valid_brand(s):
    """品牌名合法性判定"""
    if not s:
        return False
    s = s.strip()
    if len(s) > 18:
        return False
    if any(ch in s for ch in '。，；：、'):
        return False
    if re.search(r'\d{2}\.\d+\.\d+', s):  # 日期 26.9.23
        return False
    if s.startswith(('排序逻辑', '选品原则', '微唯宝', '京东', '备注', '品牌', '系列', '规格', '序号')):
        return False
    return True


def norm_sub(name):
    """规范化小类名"""
    if name is None:
        return None
    s = str(name).strip()
    s = re.split(r'——', s)[0].strip()
    s = re.sub(r'^\d+\.\s*', '', s).strip()
    s = re.sub(r'^[一二三四五六七八九十]+、', '', s).strip()
    s = s.rstrip(':：').strip()
    return s


def extract_sheet(ws, sheet_name):
    rows = []
    for r in range(1, ws.max_row + 1):
        rows.append([ws.cell(r, c).value for c in range(1, ws.max_column + 1)])

    # 找表头行 + brand_col
    header_idx = None
    brand_col = None
    for i, row in enumerate(rows):
        vals = [str(v).strip() if v is not None else '' for v in row]
        for j, v in enumerate(vals):
            if v in ('品牌', '品牌/书名'):
                header_idx = i
                brand_col = j
                break
        if header_idx is not None:
            break
    if brand_col is None:
        return []

    header = [str(v).strip() if v is not None else '' for v in rows[header_idx]]
    struct = 'A'
    if header[0] == '品类':
        struct = 'D'
    elif header[0] == '品牌':
        struct = 'C'

    # 从第 1 行开始扫描（关键修复）
    results = []
    current_sub = None
    for i in range(len(rows)):
        row = rows[i]
        v1 = row[0] if len(row) > 0 else None
        v2 = row[1] if len(row) > 1 else None

        # 跳过表头行
        if i == header_idx:
            continue
        # 跳过表头行下方的重复表头
        if isinstance(v1, str) and v1.strip() in ('序号', '品牌', '品类'):
            continue

        # 标题行检测
        if isinstance(v1, str):
            s = v1.strip()
            if '部分' in s and '第' in s:
                continue
            if re.match(r'^[一二三四五六七八九十]+、', s) and len(s) < 40:
                current_sub = norm_sub(s)
                continue
            if re.match(r'^\d+\.', s) and len(s) < 60 and (v2 is None or str(v2).strip() == ''):
                current_sub = norm_sub(s)
                continue
            if '——' in s and len(s) < 40 and (v2 is None or str(v2).strip() == ''):
                t = norm_sub(s)
                if t and len(t) < 20:
                    current_sub = t
                    continue

        # 数据行
        if struct == 'D':
            sub = str(v1).strip() if v1 is not None else None
            brand = str(row[brand_col]).strip() if len(row) > brand_col and row[brand_col] is not None else None
            if sub and brand and sub not in ('品类', '品牌', '序号') and is_valid_brand(brand):
                results.append((sub, brand))
        else:
            brand = str(row[brand_col]).strip() if len(row) > brand_col and row[brand_col] is not None else None
            if brand and is_valid_brand(brand) and not brand.startswith(('微唯宝', '京东', '备注')):
                sub = current_sub or sheet_name
                results.append((sub, brand))
    return results


all_data = {}
for f in FILES:
    fp = os.path.join(P, f)
    if not os.path.exists(fp):
        continue
    wb = openpyxl.load_workbook(fp, data_only=True)
    cat_name = f.replace('.xlsx', '')
    # 收集 (sub, brand)，保持顺序去重
    by_sub = {}
    for sn in wb.sheetnames:
        if '汇总' in sn:
            continue
        for sub, brand in extract_sheet(wb[sn], sn):
            by_sub.setdefault(sub, [])
            if brand not in by_sub[sub]:  # 去重
                by_sub[sub].append(brand)
    all_data[cat_name] = by_sub

# 输出
for cat, by_sub in all_data.items():
    print(f'\n=== {cat} ({len(by_sub)} 小类) ===')
    for sub, brands in by_sub.items():
        print(f'  {sub}: {" | ".join(brands)}')

with open(os.path.join(P, '_extracted.json'), 'w', encoding='utf-8') as fh:
    json.dump(all_data, fh, ensure_ascii=False, indent=2)
print('\n保存 _extracted.json')

# 统计
total_brands = set()
for cat, by_sub in all_data.items():
    for sub, brands in by_sub.items():
        total_brands.update(brands)
print(f'\n小类总数: {sum(len(v) for v in all_data.values())}')
print(f'唯一品牌: {len(total_brands)}')
